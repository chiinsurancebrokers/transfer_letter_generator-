import streamlit as st
from docx import Document
from io import BytesIO

st.set_page_config(page_title="Μεταφορά Παραγωγής", layout="centered")

# 🔹 Κωδικοί Brokers Union
BROKERS_UNION_CODES = {
    "AIG": "B2038",
    "Interasco": "20719",
    "Interamerican": "X770101 / X770119",
    "Allianz Ευρωπαϊκή": "Π1123",
    "Expert & GMI": "200476",
    "Υδρόγειος": "81803",
    "Generali": "13195 / 13196",
    "Generali Ι (πρώην ΑΧΑ)": "202206 / MEDISYN 84.00.00.81",
    "Intersalonica": "67690 / 7690",
    "ERGO": "7155",
    "MINETTA": "70139",
    "Eurolife": "33631 / 3631",
    "INTERLIFE": "222-14",
    "EUROINS": "1835",
    "CROMAR/LLOYD’S": "431",
    "GROUPAMA": "W36828",
    "ΑΤΛΑΝΤΙΚΗ": "1690000002",
    "ΕΘΝΙΚΗ": "49664 / 49868",
    "ΕΥΡΩΠΗ": "600642",
    "KARAVIAS": "Ρ.00383",
    "SOUTHEASTERN": "210965",
    "EUROCOVER": "664",
    "APEIRON": "11307",
    "ΣΥΝΕΤΑΙΡΙΣΤΙΚΗ": "90211",
    "BROKINS / UNIQUE": "07983",
    "ARAG": "1898/2",
    "NNHELLAS / METLIFE": "655-00-000",
    "HOWDEN": "3046",
    "HELLAS DIRECT": "BR00000000004764",
    "EXTRA ASSISTANCE": "4732",
    "ΔΥΝΑΜΙΣ": "D009560",
}

# 🔹 Defaults
DEFAULT_SUBCODE = "13041"
DEFAULT_NAME = "Χρήστος Ιατρόπουλος"

st.title("Αίτημα Μεταφοράς Παραγωγής")

# =========================
# 🔹 Παραλήπτης (Dynamic)
# =========================
st.subheader("Στοιχεία παραλήπτη")

recipient_name = st.text_input(
    "Γραφείο / Εταιρεία",
    value="CA Insurance Brokers"
)

recipient_address = st.text_input(
    "Διεύθυνση",
    value="Λεωφόρος Κηφισίας 119, Τ.Κ. 151 24, Μαρούσι, Ελλάδα"
)

recipient_phone = st.text_input(
    "Τηλέφωνο",
    value="+30 210 6100990"
)

recipient_email = st.text_input(
    "Email",
    value="info@ca-brokers.gr"
)

# =========================
# 🔹 Στοιχεία μεταφοράς
# =========================
st.subheader("Στοιχεία μεταφοράς")

company = st.selectbox("Εταιρεία", list(BROKERS_UNION_CODES.keys()))

old_company_code = st.text_input(
    "Κωδικός στην εταιρεία (συμπληρώνεται από το γραφείο)",
    value="____________"
)

brokers_union_code = st.text_input(
    "Κωδικός Brokers Union",
    value=BROKERS_UNION_CODES[company]
)

subcode = st.text_input("Υποκωδικός", value=DEFAULT_SUBCODE)
name = st.text_input("Ονοματεπώνυμο", value=DEFAULT_NAME)

# =========================
# 🔹 Κείμενο
# =========================
text = f"""Προς:
{recipient_name}
{recipient_address}
Τηλ.: {recipient_phone}
Email: {recipient_email}

Θέμα: Αίτημα μεταφοράς παραγωγής από την {company}

Παρακαλώ για τη μεταφορά της παραγωγής μου, την οποία διατηρώ στην εταιρεία {company} με κωδικό {old_company_code}, στην Brokers Union με κωδικό {brokers_union_code}, στον υποκωδικό {subcode} – {name}.

Ακόμη, έχω ενημερώσει το πελατολόγιό μου τηλεφωνικά για την αλλαγή της εξυπηρέτησής τους και έχω λάβει τη συγκατάθεσή τους.

Παρακαλώ επίσης, όπως σταλεί σε μορφή Excel η συνολική μου παραγωγή για την παρακολούθηση των μεταφορών.

Με εκτίμηση,

{name}
"""

st.subheader("Έτοιμο κείμενο")
st.text_area("Κείμενο προς αποστολή", text, height=350)

# =========================
# 🔹 Word Export
# =========================
def create_word_file(content):
    doc = Document()
    
    doc.add_heading('ΑΙΤΗΜΑ ΜΕΤΑΦΟΡΑΣ ΠΑΡΑΓΩΓΗΣ', 0)
    doc.add_paragraph(content)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

word_file = create_word_file(text)

st.download_button(
    "Κατέβασμα σε Word",
    word_file,
    file_name=f"metafora_paragogis_{company}.docx",
    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

st.download_button(
    "Κατέβασμα σε TXT",
    text,
    file_name=f"metafora_paragogis_{company}.txt",
    mime="text/plain"
)
